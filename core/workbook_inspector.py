"""
workbook_inspector.py

INPUT:
    An openpyxl Workbook (from workbook_loader.load_workbook).

PROCESS:
    Walks every sheet in the workbook and records its visibility state
    and size. This is the reusable "Excel Workbook Inspector" described
    in the project spec (section 22) - later it will grow to also count
    formula types, merged cells, defined names, etc. For Phase 1 we only
    need enough to confirm the required sheets exist.

OUTPUT:
    A pandas DataFrame, one row per sheet, columns:
    sheet_name, visibility, row_count, column_count
"""

from dataclasses import dataclass

import pandas as pd
from openpyxl.workbook.workbook import Workbook


@dataclass
class SheetSummary:
    """One row of inspection results for a single worksheet."""
    sheet_name: str
    visibility: str  # "visible" | "hidden" | "veryHidden"
    row_count: int
    column_count: int


def inspect_sheets(wb: Workbook) -> pd.DataFrame:
    """
    Summarize every sheet in the workbook: name, visible/hidden state,
    and size (rows x columns).

    Args:
        wb: an openpyxl Workbook.

    Returns:
        DataFrame with one row per sheet.
    """
    summaries: list[SheetSummary] = []
    for name in wb.sheetnames:
        ws = wb[name]
        summaries.append(
            SheetSummary(
                sheet_name=name,
                visibility=ws.sheet_state,
                row_count=ws.max_row or 0,
                column_count=ws.max_column or 0,
            )
        )
    return pd.DataFrame(summaries)


def check_required_sheets(wb: Workbook, required_sheets: list[str]) -> pd.DataFrame:
    """
    Confirm that every sheet the app depends on is actually present in
    the uploaded workbook. This is the first gate before we try to
    parse any data - if a required sheet is missing or renamed, we want
    a clear error instead of a confusing crash deeper in the pipeline.

    Args:
        wb: an openpyxl Workbook.
        required_sheets: list of sheet names the app expects to find.

    Returns:
        DataFrame with columns: sheet_name, found (bool)
    """
    present = set(wb.sheetnames)
    rows = [{"sheet_name": name, "found": name in present} for name in required_sheets]
    return pd.DataFrame(rows)
